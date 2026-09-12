import math
import os
import random
import json
import zipfile
from collections import defaultdict
from datetime import date, datetime, timedelta

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter

SEED = 20260912
rng = np.random.default_rng(SEED)
random.seed(SEED)

OUT_DIR = "/mnt/data/KPC_Synthetic_Control_Plane_Dataset_v1"
os.makedirs(OUT_DIR, exist_ok=True)

KPC_SOURCE = "https://www.kpc.co.ke/wp-content/uploads/2026/01/Kenya_Pipeline_Company_IPO-4.pdf"
EPRA_SOURCE_2025 = "https://www.epra.go.ke/sites/default/files/2025-09/Statistics-Report-June-2025-Web.pdf"
EPRA_SOURCE_2026 = "https://www.epra.go.ke/sites/default/files/2026-03/Biannual%20Statistics%20Report%202025-2026_0.pdf"
EPRA_STATS = "https://www.epra.go.ke/statistics-0"

# Official published storage capacities for inland depots (m3).
capacities = {
    ("Nairobi", "PMS"): 71726,
    ("Nairobi", "AGO"): 66551,
    ("Nairobi", "DPK"): 57187,
    ("Nairobi", "JET_A1"): 37116,
    ("Nakuru", "PMS"): 12163,
    ("Nakuru", "AGO"): 15702,
    ("Nakuru", "DPK"): 2668,
    ("Kisumu", "PMS"): 14371,
    ("Kisumu", "AGO"): 29388,
    ("Kisumu", "DPK"): 5013,
    ("Kisumu", "JET_A1"): 6516,
    ("Eldoret", "PMS"): 15471,
    ("Eldoret", "AGO"): 21922,
    ("Eldoret", "DPK"): 4413,
    ("Eldoret", "JET_A1"): 6283,
}

pipeline_cfg = {
    "Nairobi": {"line": "Line 5 (Mombasa-Nairobi)", "flow": 1000, "lead_days": 1},
    "Nakuru": {"line": "Western corridor (Line 4 scenario)", "flow": 510, "lead_days": 1},
    "Eldoret": {"line": "Line 4 (Nairobi-Eldoret)", "flow": 510, "lead_days": 2},
    "Kisumu": {"line": "Line 6 (Sinendet-Kisumu)", "flow": 290, "lead_days": 2},
}

base_loading_bays = {"Nairobi": 12, "Nakuru": 4, "Kisumu": 5, "Eldoret": 5}
road_share = {"PMS": 0.88, "AGO": 0.88, "DPK": 0.82, "JET_A1": 0.42}
tanker_m3 = 32.0

# EPRA annual domestic demand totals, used as calibration anchors. 2026 is a synthetic scenario annualization.
national_annual_m3 = {
    2021: 5_488_546.18,
    2022: 5_738_653.82,
    2023: 5_577_690.60,
    2024: 5_460_436.82,
    2025: 5_839_464.78,
    2026: 6_220_000.00,  # scenario assumption informed by 2025 H2 growth, not an official forecast
}
product_share = {"AGO": 0.47, "PMS": 0.40, "JET_A1": 0.12, "DPK": 0.01}
network_capture = {"AGO": 0.78, "PMS": 0.78, "JET_A1": 0.38, "DPK": 0.70}
depot_share = {
    "PMS": {"Nairobi": 0.58, "Nakuru": 0.11, "Kisumu": 0.14, "Eldoret": 0.17},
    "AGO": {"Nairobi": 0.50, "Nakuru": 0.10, "Kisumu": 0.20, "Eldoret": 0.20},
    "DPK": {"Nairobi": 0.65, "Nakuru": 0.06, "Kisumu": 0.15, "Eldoret": 0.14},
    "JET_A1": {"Nairobi": 0.75, "Kisumu": 0.13, "Eldoret": 0.12},
}

# EPRA FY2024/25 monthly totals normalized to create an official-data-informed seasonality profile.
monthly_total_anchor = {
    1: 505.77, 2: 466.80, 3: 491.13, 4: 497.94, 5: 495.98, 6: 470.63,
    7: 489.01, 8: 487.26, 9: 472.89, 10: 498.61, 11: 467.09, 12: 496.35,
}
month_mean = np.mean(list(monthly_total_anchor.values()))
monthly_factor = {m: v / month_mean for m, v in monthly_total_anchor.items()}

dow_factor = {0: 1.02, 1: 1.03, 2: 1.03, 3: 1.04, 4: 1.08, 5: 0.94, 6: 0.86}

# Scenario reference prices - explicitly synthetic, not historical EPRA monthly pump prices.
price_year_base = {
    2021: {"PMS": 130, "AGO": 112, "DPK": 105, "JET_A1": 92},
    2022: {"PMS": 160, "AGO": 142, "DPK": 130, "JET_A1": 120},
    2023: {"PMS": 190, "AGO": 185, "DPK": 175, "JET_A1": 165},
    2024: {"PMS": 194, "AGO": 180, "DPK": 170, "JET_A1": 160},
    2025: {"PMS": 178, "AGO": 166, "DPK": 151, "JET_A1": 148},
    2026: {"PMS": 207, "AGO": 210, "DPK": 184, "JET_A1": 176},
}
depot_price_adder = {"Nairobi": 3.0, "Nakuru": 2.2, "Kisumu": 3.0, "Eldoret": 3.0}


def easter_sunday(year):
    # Anonymous Gregorian algorithm
    a = year % 19
    b = year // 100
    c = year % 100
    d = b // 4
    e = b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i = c // 4
    k = c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    month = (h + l - 7 * m + 114) // 31
    day = ((h + l - 7 * m + 114) % 31) + 1
    return date(year, month, day)


def holiday_dates(year):
    es = easter_sunday(year)
    fixed = {
        date(year, 1, 1): "New Year",
        date(year, 5, 1): "Labour Day",
        date(year, 6, 1): "Madaraka Day",
        date(year, 10, 20): "Mashujaa Day",
        date(year, 12, 12): "Jamhuri Day",
        date(year, 12, 25): "Christmas Day",
        date(year, 12, 26): "Boxing Day",
        es - timedelta(days=2): "Good Friday",
        es + timedelta(days=1): "Easter Monday",
    }
    return fixed

all_holidays = {}
for y in range(2021, 2027):
    all_holidays.update(holiday_dates(y))

# Generate equipment events first so both baseline/autonomous policies see identical operational disruptions.
event_rows = []
event_map = defaultdict(lambda: {"pipeline_available": 1, "pump_available": 1, "bay_reduction": 0, "event_labels": []})
for depot in ["Nairobi", "Nakuru", "Kisumu", "Eldoret"]:
    for year in range(2021, 2027):
        year_start = date(year, 1, 10)
        year_end = date(year, 12, 15)
        span = (year_end - year_start).days
        specs = [
            ("PLANNED_PIPELINE_MAINTENANCE", 2, (1, 2), "MEDIUM", 100),
            ("UNPLANNED_PUMP_FAILURE", 3, (1, 1), "HIGH", 100),
            ("LOADING_BAY_OUTAGE", 5, (1, 1), "MEDIUM", 35),
            ("DATA_LATENCY_INCIDENT", 2, (1, 1), "LOW", 0),
        ]
        for etype, count, dur_range, severity, impact in specs:
            for _ in range(count):
                start = year_start + timedelta(days=int(rng.integers(0, span + 1)))
                duration = int(rng.integers(dur_range[0], dur_range[1] + 1))
                end = start + timedelta(days=duration - 1)
                event_id = f"EV-{depot[:3].upper()}-{year}-{len(event_rows)+1:05d}"
                event_rows.append({
                    "event_id": event_id,
                    "depot": depot,
                    "event_type": etype,
                    "severity": severity,
                    "start_date": start,
                    "end_date": end,
                    "duration_days": duration,
                    "impact_percent": impact,
                    "source_type": "SYNTHETIC_SCENARIO",
                })
                for i in range(duration):
                    d = start + timedelta(days=i)
                    em = event_map[(depot, d)]
                    em["event_labels"].append(etype)
                    if etype == "PLANNED_PIPELINE_MAINTENANCE":
                        em["pipeline_available"] = 0
                    elif etype == "UNPLANNED_PUMP_FAILURE":
                        em["pump_available"] = 0
                        em["pipeline_available"] = 0
                    elif etype == "LOADING_BAY_OUTAGE":
                        em["bay_reduction"] = max(em["bay_reduction"], 2 if depot == "Nairobi" else 1)

equipment_events = pd.DataFrame(event_rows)

# Generate demand/external features for all valid depot-product combinations.
start = date(2021, 1, 1)
end = date(2026, 12, 31)
all_dates = pd.date_range(start, end, freq="D")

feature_rows = []
for (depot, product), capacity in capacities.items():
    for year in range(2021, 2027):
        dates_y = [d.date() for d in all_dates if d.year == year]
        target = national_annual_m3[year] * product_share[product] * network_capture[product] * depot_share[product][depot]
        raw_weights = []
        for d in dates_y:
            f = monthly_factor[d.month] * dow_factor[d.weekday()]
            # Product-specific seasonal behavior.
            if product == "PMS" and (d.month == 12 or (d.month == 1 and d.day <= 5)):
                f *= 1.14
            if product == "AGO" and d.month == 10:
                f *= 1.10
            if product == "AGO" and depot in ("Eldoret", "Kisumu") and d.month in (9, 10, 11):
                f *= 1.07
            if product == "JET_A1" and d.month == 12:
                f *= 1.10
            if product == "DPK" and d.month in (6, 7):
                f *= 1.12
            # Holidays and school reopening demand effects.
            if d in all_holidays:
                if product == "PMS":
                    f *= 1.08
                elif product == "JET_A1":
                    f *= 1.06
            if (d.month == 1 and 3 <= d.day <= 10) or (d.month == 5 and 1 <= d.day <= 7) or (d.month == 8 and 24 <= d.day <= 31):
                if product in ("PMS", "AGO"):
                    f *= 1.04
            # Random daily demand noise; small enough to preserve operational realism.
            f *= float(rng.lognormal(mean=0, sigma=0.055))
            raw_weights.append(f)
        raw_weights = np.array(raw_weights)
        demands = target * raw_weights / raw_weights.sum()
        for d, demand_m3 in zip(dates_y, demands):
            # Heavy rain seasons: Mar-May and Oct-Dec.
            seasonal_rain = 1.8 if d.month in (3, 4, 5, 10, 11, 12) else 0.8
            rainfall = float(rng.gamma(shape=1.4, scale=seasonal_rain))
            # Depot-specific rainfall adjustment.
            rainfall *= {"Nairobi": 1.0, "Nakuru": 0.9, "Kisumu": 1.35, "Eldoret": 1.1}[depot]
            price_base = price_year_base[year][product] + depot_price_adder[depot]
            month_wave = 1.0 + 0.025 * math.sin((d.month - 1) / 12 * 2 * math.pi)
            price = price_base * month_wave + float(rng.normal(0, 1.0))
            holiday_name = all_holidays.get(d, "")
            event = event_map[(depot, d)]
            feature_rows.append({
                "date": d,
                "depot": depot,
                "product": product,
                "tank_capacity_m3": float(capacity),
                "customer_orders_m3": float(demand_m3),
                "holiday_flag": 1 if holiday_name else 0,
                "holiday_name": holiday_name,
                "rainfall_mm": max(0.0, rainfall),
                "reference_price_kes_l": max(1.0, price),
                "pipeline_available": event["pipeline_available"],
                "pump_available": event["pump_available"],
                "bay_reduction": event["bay_reduction"],
                "event_label": ";".join(sorted(set(event["event_labels"]))) if event["event_labels"] else "NORMAL",
            })

base_features = pd.DataFrame(feature_rows)
base_features.sort_values(["depot", "product", "date"], inplace=True)

# Add a few deterministic demand spikes to make anomaly detection and control-plane behavior visible.
spike_specs = [
    (date(2023, 12, 22), "Nairobi", "PMS", 1.28),
    (date(2024, 10, 4), "Eldoret", "AGO", 1.34),
    (date(2025, 4, 17), "Kisumu", "PMS", 1.24),
    (date(2025, 12, 19), "Nairobi", "JET_A1", 1.32),
    (date(2026, 8, 28), "Eldoret", "AGO", 1.35),
    (date(2026, 9, 9), "Kisumu", "AGO", 1.29),
    (date(2026, 12, 20), "Nairobi", "PMS", 1.31),
]
for d, depot, product, mult in spike_specs:
    mask = (base_features["date"] == d) & (base_features["depot"] == depot) & (base_features["product"] == product)
    base_features.loc[mask, "customer_orders_m3"] *= mult
    base_features.loc[mask, "event_label"] = base_features.loc[mask, "event_label"].apply(lambda x: (x + ";DEMAND_SPIKE").strip(";"))


def simulate_policy(name, policy_cfg):
    ops_rows = []
    batch_rows = []
    pending = defaultdict(list)  # (depot, product) -> batch dicts not yet arrived
    batch_counter = 0
    last_close = {}

    for (depot, product), grp in base_features.groupby(["depot", "product"], sort=True):
        grp = grp.sort_values("date")
        cap = capacities[(depot, product)]
        safety_frac = {"PMS": 0.27, "AGO": 0.28, "DPK": 0.22, "JET_A1": 0.22}[product]
        target_frac = policy_cfg["target_frac"]
        reorder_frac = policy_cfg["reorder_frac"]
        base_lead = pipeline_cfg[depot]["lead_days"] + policy_cfg.get("lead_days_extra", 0)
        max_flow = pipeline_cfg[depot]["flow"]
        line = pipeline_cfg[depot]["line"]
        base_bays = base_loading_bays[depot]
        # initial stock between 55-68% of capacity.
        opening = cap * float(rng.uniform(0.55, 0.68))

        for rec in grp.itertuples(index=False):
            d = rec.date
            orders = float(rec.customer_orders_m3)
            # Arriving pipeline batches for the day.
            arrivals = [b for b in pending[(depot, product)] if b["actual_arrival_date"] == d]
            pipeline_receipts = 0.0
            if arrivals:
                requested_arrival = sum(b["scheduled_volume_m3"] for b in arrivals)
                pipeline_receipts = min(requested_arrival, max(0.0, cap - opening))
                remaining = pipeline_receipts
                for b in arrivals:
                    delivered = min(b["scheduled_volume_m3"], remaining)
                    b["delivered_volume_m3"] = delivered
                    remaining -= delivered
                    b["status"] = "COMPLETED" if d <= date(2026, 9, 12) else "SCENARIO_COMPLETED"
            # Remove arrived batches from pending.
            pending[(depot, product)] = [b for b in pending[(depot, product)] if b["actual_arrival_date"] > d]

            bays = max(1, base_bays - int(rec.bay_reduction))
            loading_capacity = bays * 25 * tanker_m3
            inventory_before_dispatch = opening + pipeline_receipts

            # Reactive emergency road transfer is allowed; autonomous policy uses it earlier.
            emergency_receipts = 0.0
            projected_shortfall = max(0.0, orders - min(inventory_before_dispatch, loading_capacity))
            if policy_cfg["emergency_enabled"]:
                emergency_trigger = cap * policy_cfg["emergency_trigger_frac"]
                if inventory_before_dispatch < emergency_trigger or projected_shortfall > 0:
                    desired_buffer = cap * policy_cfg["emergency_target_frac"]
                    emergency_receipts = min(
                        max(0.0, desired_buffer - inventory_before_dispatch),
                        max(700.0, cap * 0.09),
                        max(0.0, cap - inventory_before_dispatch),
                    )
                    inventory_before_dispatch += emergency_receipts

            dispatch = min(orders, inventory_before_dispatch, loading_capacity)
            unfulfilled = max(0.0, orders - dispatch)
            closing = max(0.0, inventory_before_dispatch - dispatch)
            safety_stock = cap * safety_frac
            reorder_point = max(cap * reorder_frac, safety_stock)
            reorder_point = min(reorder_point, cap * 0.60)
            days_cover = closing / max(orders, 1.0)

            road_vol = dispatch * road_share[product]
            trucks_requested = int(math.ceil(orders * road_share[product] / tanker_m3))
            trucks_loaded = int(math.ceil(road_vol / tanker_m3)) if road_vol > 0 else 0
            utilization = trucks_loaded / max(1, bays * 25)
            turnaround = 50 + 95 * (utilization ** 2) + (45 if rec.bay_reduction else 0) + float(rng.normal(0, 7))
            turnaround = max(35.0, min(260.0, turnaround))
            queue = max(0.0, turnaround - 42.0)
            delayed_rate = min(0.65, max(0.02, (turnaround - 55) / 220))
            delayed_trucks = int(round(trucks_loaded * delayed_rate))

            # Inventory position includes receipts already planned over the next few days.
            pending_volume = sum(b["scheduled_volume_m3"] for b in pending[(depot, product)])
            inventory_position = closing + pending_volume
            batch_triggered = 0
            batch_volume = 0.0
            action = "NONE"

            if inventory_position < reorder_point:
                batch_triggered = 1
                forecast_during_lead = orders * (base_lead + 0.5)
                desired_arrival_stock = cap * target_frac
                batch_volume = max(0.0, desired_arrival_stock + forecast_during_lead - inventory_position)
                daily_corridor_limit = max_flow * policy_cfg.get("pipeline_operating_hours", 16)
                batch_volume = min(batch_volume, daily_corridor_limit, cap * 0.48)
                # Avoid tiny operational batches.
                batch_volume = max(batch_volume, min(cap * 0.12, 1200.0))
                # Determine delay. If planned arrival hits an outage, push by one day until available.
                scheduled_arrival = d + timedelta(days=base_lead)
                delay_days = 0
                r = float(rng.random())
                if r < policy_cfg["batch_delay_prob"]:
                    delay_days = 1 if r > policy_cfg["batch_delay_prob"] * 0.15 else 2
                actual_arrival = scheduled_arrival + timedelta(days=delay_days)
                while event_map[(depot, actual_arrival)]["pipeline_available"] == 0 or event_map[(depot, actual_arrival)]["pump_available"] == 0:
                    actual_arrival += timedelta(days=1)
                    delay_days += 1
                batch_counter += 1
                batch_id = f"{name[:3].upper()}-B-{batch_counter:06d}"
                b = {
                    "batch_id": batch_id,
                    "policy": name,
                    "trigger_date": d,
                    "origin": "Mombasa/Kipevu" if depot == "Nairobi" else "Nairobi/Upstream Hub",
                    "destination": depot,
                    "product": product,
                    "pipeline_line": line,
                    "flow_limit_m3_hr": max_flow,
                    "scheduled_volume_m3": batch_volume,
                    "scheduled_arrival_date": scheduled_arrival,
                    "actual_arrival_date": actual_arrival,
                    "delay_days": delay_days,
                    "delivered_volume_m3": 0.0,
                    "reason": "REORDER_POINT_BREACH",
                    "automated_trigger": 1 if name == "AUTONOMOUS" else 0,
                    "status": "PLANNED" if actual_arrival > date(2026, 9, 12) else "PENDING",
                }
                pending[(depot, product)].append(b)
                batch_rows.append(b)
                action = "CREATE_PIPELINE_BATCH" if name == "AUTONOMOUS" else "MANUAL_REPLENISHMENT_ORDER"

            if emergency_receipts > 0:
                action = "EMERGENCY_TRUCK_TRANSFER" if action == "NONE" else action + "+EMERGENCY_TRUCK_TRANSFER"
            if rec.pipeline_available == 0 or rec.pump_available == 0:
                action = "HOLD/RESCHEDULE_DUE_TO_OUTAGE" if action == "NONE" else action + "+OUTAGE"

            if unfulfilled > 0:
                stock_status = "STOCKOUT"
            elif closing < safety_stock:
                stock_status = "CRITICAL"
            elif closing < reorder_point:
                stock_status = "WATCH"
            else:
                stock_status = "HEALTHY"

            record_status = "SIMULATED_HISTORY" if d <= date(2026, 9, 12) else "SCENARIO_FUTURE"
            capacity_breach = int(closing > cap + 1e-6 or opening > cap + 1e-6)
            reconciliation = abs(closing - (opening + pipeline_receipts + emergency_receipts - dispatch))
            recon_err = int(reconciliation > 0.01)

            ops_rows.append({
                "date": d,
                "record_status": record_status,
                "policy": name,
                "depot": depot,
                "product": product,
                "tank_capacity_m3": round(cap, 2),
                "opening_stock_m3": round(opening, 2),
                "customer_orders_m3": round(orders, 2),
                "pipeline_receipts_m3": round(pipeline_receipts, 2),
                "emergency_truck_receipts_m3": round(emergency_receipts, 2),
                "dispatch_m3": round(dispatch, 2),
                "unfulfilled_demand_m3": round(unfulfilled, 2),
                "closing_stock_m3": round(closing, 2),
                "safety_stock_m3": round(safety_stock, 2),
                "reorder_point_m3": round(reorder_point, 2),
                "days_of_cover": round(days_cover, 2),
                "pipeline_line": line,
                "pipeline_flow_limit_m3_hr": max_flow,
                "pipeline_available": int(rec.pipeline_available),
                "pump_available": int(rec.pump_available),
                "loading_bays_available": bays,
                "loading_capacity_m3": round(loading_capacity, 2),
                "trucks_requested": trucks_requested,
                "trucks_loaded": trucks_loaded,
                "avg_truck_turnaround_min": round(turnaround, 1),
                "queue_time_min": round(queue, 1),
                "delayed_trucks": delayed_trucks,
                "batch_triggered": batch_triggered,
                "batch_volume_m3": round(batch_volume, 2),
                "emergency_replenishment": int(emergency_receipts > 0),
                "holiday_flag": int(rec.holiday_flag),
                "event_label": rec.event_label,
                "rainfall_mm": round(float(rec.rainfall_mm), 2),
                "reference_price_kes_l": round(float(rec.reference_price_kes_l), 2),
                "day_of_week": d.strftime("%A"),
                "month": d.month,
                "year": d.year,
                "stock_status": stock_status,
                "autonomous_action": action,
                "capacity_breach_flag": capacity_breach,
                "inventory_reconciliation_error": recon_err,
            })
            opening = closing

    ops = pd.DataFrame(ops_rows)
    batches = pd.DataFrame(batch_rows)
    return ops, batches

# Autonomous control plane policy vs a slower reactive baseline policy.
autonomous_cfg = {
    "target_frac": 0.76,
    "reorder_frac": 0.43,
    "lead_days_extra": 0,
    "emergency_enabled": True,
    "emergency_trigger_frac": 0.28,
    "emergency_target_frac": 0.36,
    "batch_delay_prob": 0.07,
    "pipeline_operating_hours": 16,
}
baseline_cfg = {
    "target_frac": 0.45,
    "reorder_frac": 0.12,
    "lead_days_extra": 3,
    "emergency_enabled": True,
    "emergency_trigger_frac": 0.06,
    "emergency_target_frac": 0.10,
    "batch_delay_prob": 0.15,
    "pipeline_operating_hours": 10,
}

auto_ops, auto_batches = simulate_policy("AUTONOMOUS", autonomous_cfg)
base_ops, base_batches = simulate_policy("BASELINE", baseline_cfg)

# Use autonomous policy as the primary operations dataset.
daily_ops = auto_ops.copy()

# Customer orders sample for final 120 days, split across synthetic OMCs for demo/API payload testing.
omcs = ["OMC_A", "OMC_B", "OMC_C", "OMC_D", "OMC_E", "OMC_F"]
customer_order_rows = []
recent = daily_ops[daily_ops["date"] >= date(2026, 9, 3)].copy()
order_id = 0
for r in recent.itertuples(index=False):
    shares = rng.dirichlet(np.array([3.0, 2.5, 2.0, 1.7, 1.2, 0.9]))
    for omc, sh in zip(omcs, shares):
        order_id += 1
        q = r.customer_orders_m3 * sh
        fulfilled = r.dispatch_m3 * sh
        customer_order_rows.append({
            "order_id": f"ORD-{order_id:07d}",
            "date": r.date,
            "depot": r.depot,
            "product": r.product,
            "customer_id": omc,
            "ordered_m3": round(q, 2),
            "fulfilled_m3": round(min(q, fulfilled), 2),
            "priority": "HIGH" if float(rng.random()) < 0.08 else "NORMAL",
            "source_type": "SYNTHETIC_SCENARIO",
        })
customer_orders = pd.DataFrame(customer_order_rows)

# Aggregated truck operations convenience table.
truck_ops = daily_ops[[
    "date", "record_status", "depot", "product", "trucks_requested", "trucks_loaded",
    "avg_truck_turnaround_min", "queue_time_min", "delayed_trucks", "dispatch_m3",
    "loading_bays_available", "loading_capacity_m3", "event_label"
]].copy()

# Model-ready forecasting features with no forward-looking demand leakage.
forecast = daily_ops[[
    "date", "record_status", "depot", "product", "customer_orders_m3", "holiday_flag",
    "rainfall_mm", "reference_price_kes_l", "day_of_week", "month", "year"
]].copy()
forecast.sort_values(["depot", "product", "date"], inplace=True)
for lag in [1, 7, 14, 28]:
    forecast[f"demand_lag_{lag}_m3"] = forecast.groupby(["depot", "product"])["customer_orders_m3"].shift(lag)
forecast["rolling_7d_mean_m3"] = forecast.groupby(["depot", "product"])["customer_orders_m3"].transform(lambda s: s.shift(1).rolling(7, min_periods=3).mean())
forecast["rolling_28d_mean_m3"] = forecast.groupby(["depot", "product"])["customer_orders_m3"].transform(lambda s: s.shift(1).rolling(28, min_periods=7).mean())
forecast["dow_num"] = pd.to_datetime(forecast["date"]).dt.dayofweek
forecast["dow_sin"] = np.sin(2 * np.pi * forecast["dow_num"] / 7)
forecast["dow_cos"] = np.cos(2 * np.pi * forecast["dow_num"] / 7)
forecast["month_sin"] = np.sin(2 * np.pi * (forecast["month"] - 1) / 12)
forecast["month_cos"] = np.cos(2 * np.pi * (forecast["month"] - 1) / 12)
forecast["baseline_forecast_m3"] = 0.55 * forecast["demand_lag_7_m3"] + 0.45 * forecast["rolling_28d_mean_m3"]
forecast = forecast.round({
    "customer_orders_m3": 2, "rainfall_mm": 2, "reference_price_kes_l": 2,
    "demand_lag_1_m3": 2, "demand_lag_7_m3": 2, "demand_lag_14_m3": 2, "demand_lag_28_m3": 2,
    "rolling_7d_mean_m3": 2, "rolling_28d_mean_m3": 2, "baseline_forecast_m3": 2,
    "dow_sin": 6, "dow_cos": 6, "month_sin": 6, "month_cos": 6,
})

# Control-plane events: actionable events from autonomous dataset.
cp = daily_ops[(daily_ops["batch_triggered"] == 1) | (daily_ops["emergency_replenishment"] == 1) | (daily_ops["stock_status"].isin(["CRITICAL", "STOCKOUT"])) | (daily_ops["event_label"] != "NORMAL")].copy()
cp["control_event_id"] = [f"CP-{i+1:07d}" for i in range(len(cp))]
control_plane_events = cp[[
    "control_event_id", "date", "record_status", "depot", "product", "stock_status", "event_label",
    "opening_stock_m3", "customer_orders_m3", "closing_stock_m3", "days_of_cover",
    "batch_triggered", "batch_volume_m3", "emergency_replenishment", "autonomous_action"
]].copy()

# ROI comparison for 2026 scenario, comparing baseline vs autonomous under same generated demand/events.
def monthly_metrics(df, prefix):
    t = df[df["year"] == 2026].copy()
    t["month_key"] = pd.to_datetime(t["date"]).dt.to_period("M").astype(str)
    g = t.groupby("month_key").agg(
        unfulfilled_m3=("unfulfilled_demand_m3", "sum"),
        emergency_m3=("emergency_truck_receipts_m3", "sum"),
        batch_actions=("batch_triggered", "sum"),
        stockout_days=("stock_status", lambda s: int((s == "STOCKOUT").sum())),
        demand_m3=("customer_orders_m3", "sum"),
        dispatch_m3=("dispatch_m3", "sum"),
    ).reset_index()
    return g.rename(columns={c: f"{prefix}_{c}" for c in g.columns if c != "month_key"})

roi_compare = monthly_metrics(base_ops, "baseline").merge(monthly_metrics(auto_ops, "autonomous"), on="month_key")
roi_compare["stockout_m3_avoided"] = (roi_compare["baseline_unfulfilled_m3"] - roi_compare["autonomous_unfulfilled_m3"]).clip(lower=0)
roi_compare["emergency_m3_avoided"] = (roi_compare["baseline_emergency_m3"] - roi_compare["autonomous_emergency_m3"]).clip(lower=0)
roi_compare["service_level_baseline_pct"] = 1 - roi_compare["baseline_unfulfilled_m3"] / roi_compare["baseline_demand_m3"]
roi_compare["service_level_autonomous_pct"] = 1 - roi_compare["autonomous_unfulfilled_m3"] / roi_compare["autonomous_demand_m3"]
roi_compare = roi_compare.round(2)

# Official capacity table and calibration anchors.
cap_rows = []
for (depot, product), cap in capacities.items():
    cap_rows.append({
        "depot": depot,
        "product": product,
        "capacity_m3": cap,
        "source_type": "OFFICIAL_KPC_PUBLISHED",
        "source_url": KPC_SOURCE,
    })
depot_capacities = pd.DataFrame(cap_rows)

calibration_rows = []
for y, total in national_annual_m3.items():
    calibration_rows.append({
        "year": y,
        "national_petroleum_demand_m3": total,
        "source_type": "OFFICIAL_EPRA" if y <= 2025 else "SYNTHETIC_SCENARIO",
        "source_url": EPRA_SOURCE_2025 if y <= 2025 else EPRA_SOURCE_2026,
        "note": "EPRA annual total used as anchor" if y <= 2025 else "Scenario annualization; not an official EPRA forecast",
    })
calibration_anchors = pd.DataFrame(calibration_rows)

cost_parameters = pd.DataFrame([
    ["stockout_value_at_risk_per_m3_kes", 12000, "KES/m3", "SYNTHETIC_SCENARIO", "Adjustable ROI assumption; not a KPC disclosed cost"],
    ["emergency_logistics_premium_per_m3_kes", 450, "KES/m3", "SYNTHETIC_SCENARIO", "Premium used for emergency road transfer scenario"],
    ["manual_intervention_cost_per_action_kes", 15000, "KES/action", "SYNTHETIC_SCENARIO", "Time/coordination cost proxy"],
    ["implementation_cost_kes", 8500000, "KES", "SYNTHETIC_SCENARIO", "One-time platform implementation assumption"],
    ["annual_support_cost_kes", 1800000, "KES/year", "SYNTHETIC_SCENARIO", "Annual cloud/support/monitoring assumption"],
], columns=["parameter", "value", "unit", "source_type", "note"])

# Data dictionary.
dictionary = {
    "date": "Operating date.",
    "record_status": "SIMULATED_HISTORY through 2026-09-12; SCENARIO_FUTURE afterwards.",
    "policy": "Inventory control policy used in simulation.",
    "depot": "KPC inland depot used by the synthetic network.",
    "product": "PMS, AGO, DPK or JET_A1.",
    "tank_capacity_m3": "Published KPC product storage capacity for the depot.",
    "opening_stock_m3": "Synthetic start-of-day usable inventory.",
    "customer_orders_m3": "Synthetic requested daily demand; recommended forecasting target.",
    "pipeline_receipts_m3": "Synthetic pipeline receipt volume delivered that day.",
    "emergency_truck_receipts_m3": "Synthetic emergency inter-depot road transfer received.",
    "dispatch_m3": "Fulfilled/loaded demand volume; used in inventory reconciliation.",
    "unfulfilled_demand_m3": "Orders not fulfilled due to inventory or loading constraints.",
    "closing_stock_m3": "End-of-day stock = opening + receipts - dispatch.",
    "safety_stock_m3": "Synthetic minimum safety inventory threshold.",
    "reorder_point_m3": "Synthetic autonomous replenishment trigger threshold.",
    "days_of_cover": "Closing stock divided by the day's requested demand.",
    "pipeline_line": "Published corridor name where available; Nakuru mapping is a scenario attribution.",
    "pipeline_flow_limit_m3_hr": "Published KPC installed flow capacity for the assigned corridor.",
    "pipeline_available": "Synthetic daily availability flag (1/0).",
    "pump_available": "Synthetic pump availability flag (1/0).",
    "loading_bays_available": "Synthetic available loading-bay count.",
    "loading_capacity_m3": "Synthetic daily loading throughput capacity.",
    "trucks_requested": "Synthetic trucks implied by customer demand and road-dispatch share.",
    "trucks_loaded": "Synthetic trucks loaded based on fulfilled dispatch.",
    "avg_truck_turnaround_min": "Synthetic average truck turnaround time.",
    "queue_time_min": "Synthetic average queue component of truck turnaround.",
    "delayed_trucks": "Synthetic count of delayed trucks.",
    "batch_triggered": "1 when replenishment policy creates a pipeline batch.",
    "batch_volume_m3": "Synthetic replenishment batch quantity requested.",
    "emergency_replenishment": "1 when emergency road transfer is invoked.",
    "holiday_flag": "1 for selected Kenyan public holidays used in the scenario.",
    "event_label": "Operational anomaly/event label; NORMAL if none.",
    "rainfall_mm": "Synthetic rainfall feature.",
    "reference_price_kes_l": "Synthetic scenario reference price; not official historical pump price.",
    "day_of_week": "Calendar weekday.",
    "month": "Calendar month number.",
    "year": "Calendar year.",
    "stock_status": "HEALTHY, WATCH, CRITICAL or STOCKOUT.",
    "autonomous_action": "Action selected by the simulated control-plane policy.",
    "capacity_breach_flag": "Validation flag: 1 if stock exceeds capacity.",
    "inventory_reconciliation_error": "Validation flag: 1 if stock equation fails tolerance.",
}
data_dictionary = pd.DataFrame([{"field": k, "definition": v} for k, v in dictionary.items()])

# Save CSV files.
files_to_save = {
    "daily_operations.csv": daily_ops,
    "pipeline_batches.csv": auto_batches,
    "truck_operations_daily.csv": truck_ops,
    "customer_orders_demo.csv": customer_orders,
    "equipment_events.csv": equipment_events,
    "forecasting_features.csv": forecast,
    "control_plane_events.csv": control_plane_events,
    "depot_capacities.csv": depot_capacities,
    "calibration_anchors.csv": calibration_anchors,
    "roi_comparison_2026.csv": roi_compare,
    "cost_parameters.csv": cost_parameters,
    "data_dictionary.csv": data_dictionary,
}
for fname, df in files_to_save.items():
    df.to_csv(os.path.join(OUT_DIR, fname), index=False)

# Dataset card and README.
dataset_card = f"""# KPC Synthetic Downstream Operations Dataset v1.0

## Purpose
This dataset supports the KPC Inuka Fellowship Hackathon 3 use case: downstream demand forecasting, inventory replenishment, autonomous batch/dispatch decisions, monitoring, and executive ROI demonstration.

## Important disclosure
This is a **synthetic scenario dataset**. It does not contain confidential or actual KPC operational transaction records. Published KPC/EPRA statistics are used only as calibration anchors and hard constraints.

## Coverage
- Dates: 2021-01-01 to 2026-12-31
- Simulated-history cutoff: 2026-09-12
- Dates after 2026-09-12 are explicitly labelled `SCENARIO_FUTURE`
- Depots: Nairobi, Nakuru, Kisumu, Eldoret
- Products: PMS, AGO, DPK, JET_A1 where KPC publishes storage capacity
- Random seed: {SEED}

## Official calibration anchors
1. KPC published depot/product storage capacities and pipeline flow rates: {KPC_SOURCE}
2. EPRA annual domestic petroleum demand totals and FY2024/25 monthly demand seasonality: {EPRA_SOURCE_2025}
3. EPRA 2025/26 biannual statistics used only to inform the 2026 synthetic scenario: {EPRA_SOURCE_2026}

## Core integrity constraints
- Closing stock = opening stock + pipeline receipts + emergency receipts - dispatch
- Stock cannot exceed published tank capacity
- Unsupported product/depot combinations are excluded
- Replenishment batches are constrained by a corridor flow-rate limit
- Operational failures, demand spikes, orders, rainfall and costs are synthetic

## Recommended ML target
`customer_orders_m3`

## Recommended train/test split
- Train: 2021-01-01 to 2024-12-31
- Validation: 2025-01-01 to 2025-12-31
- Test/demo: 2026-01-01 to 2026-09-12
- Scenario-only future: after 2026-09-12

## Limitations
This dataset is intended for prototyping, hackathon evaluation and architecture demonstrations. It must not be represented as actual KPC operations or used for real operational decisions without replacement by governed production data.
"""
with open(os.path.join(OUT_DIR, "DATASET_CARD.md"), "w", encoding="utf-8") as f:
    f.write(dataset_card)

readme = """# KPC Autonomous Supply Chain & Executive Control Plane - Synthetic Data Package

Start with `daily_operations.csv` for the operational control-plane model and `forecasting_features.csv` for demand forecasting.

Useful files:
- `daily_operations.csv`: primary autonomous-policy operational dataset
- `forecasting_features.csv`: lagged/model-ready demand features
- `pipeline_batches.csv`: autonomous replenishment batch events
- `control_plane_events.csv`: actionable alerts/actions for API/demo streaming
- `roi_comparison_2026.csv`: baseline vs autonomous scenario metrics
- `KPC_Synthetic_Control_Plane_Data.xlsx`: formatted workbook with documentation, validation and ROI formulas
- `generate_kpc_dataset.py`: reproducibility script

All operational records are synthetic. See `DATASET_CARD.md` before using the data.
"""
with open(os.path.join(OUT_DIR, "README.md"), "w", encoding="utf-8") as f:
    f.write(readme)

# Copy reproducibility script into package.
import shutil
shutil.copy2("/mnt/data/generate_kpc_dataset.py", os.path.join(OUT_DIR, "generate_kpc_dataset.py"))

# Build formatted Excel workbook.
xlsx_path = os.path.join(OUT_DIR, "KPC_Synthetic_Control_Plane_Data.xlsx")
wb = Workbook()
ws = wb.active
ws.title = "README"
ws.sheet_view.showGridLines = False

# Theme colors.
DARK = "16365C"
HEADER = "1F4E78"
TEAL = "DDEBF7"
LIGHT_BLUE = "D9EAF7"
GRAY = "E7E6E6"
GREEN = "E2F0D9"
ORANGE = "FCE4D6"
RED = "F4CCCC"
PURPLE = "E4DFEC"
WHITE = "FFFFFF"
BLACK = "000000"
BLUE_TEXT = "0000FF"
GREEN_TEXT = "008000"
GRAY_TEXT = "666666"
ORANGE_TEXT = "C65911"
PURPLE_TEXT = "7030A0"
TEAL_TEXT = "008C95"

ws.merge_cells("A1:H2")
ws["A1"] = "KPC Synthetic Downstream Operations Dataset v1.0"
ws["A1"].font = Font(size=18, bold=True, color=WHITE)
ws["A1"].fill = PatternFill("solid", fgColor=DARK)
ws["A1"].alignment = Alignment(vertical="center", horizontal="left")
ws.row_dimensions[1].height = 26
ws.row_dimensions[2].height = 20

readme_rows = [
    ("Purpose", "Hackathon 3 demand forecasting, replenishment automation, monitoring and ROI demonstration."),
    ("Disclosure", "Synthetic scenario dataset. No confidential/actual KPC transaction records are included."),
    ("Coverage", "2021-01-01 to 2026-12-31; dates after 2026-09-12 are labelled SCENARIO_FUTURE."),
    ("Depots", "Nairobi, Nakuru, Kisumu, Eldoret."),
    ("Products", "PMS, AGO, DPK, JET_A1 where KPC publishes a storage capacity."),
    ("Recommended target", "customer_orders_m3"),
    ("Seed", str(SEED)),
    ("Official KPC source", KPC_SOURCE),
    ("Official EPRA source", EPRA_SOURCE_2025),
    ("2026 scenario context", EPRA_SOURCE_2026),
]
for i, (k, v) in enumerate(readme_rows, start=4):
    ws[f"A{i}"] = k
    ws[f"A{i}"].font = Font(bold=True, color=GRAY_TEXT)
    ws[f"B{i}"] = v
    ws[f"B{i}"].alignment = Alignment(wrap_text=True)
    if "source" in k.lower() or "context" in k.lower():
        ws[f"B{i}"].font = Font(color=GREEN_TEXT, underline="single")
ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 95

# Helper to add DataFrame as a styled sheet.
def add_df_sheet(name, df, table_name=None, max_rows=None):
    sh = wb.create_sheet(name)
    sh.sheet_view.showGridLines = False
    data = df if max_rows is None else df.head(max_rows)
    # header
    for c, col in enumerate(data.columns, start=1):
        cell = sh.cell(row=1, column=c, value=col)
        cell.fill = PatternFill("solid", fgColor=HEADER)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # data
    for r_idx, row in enumerate(data.itertuples(index=False, name=None), start=2):
        for c_idx, val in enumerate(row, start=1):
            if isinstance(val, (pd.Timestamp, datetime)):
                val = val.date()
            cell = sh.cell(row=r_idx, column=c_idx, value=val)
            if isinstance(val, date):
                cell.number_format = "yyyy-mm-dd"
    sh.freeze_panes = "A2"
    sh.auto_filter.ref = sh.dimensions
    sh.row_dimensions[1].height = 32
    # widths based on header.
    for idx, col in enumerate(data.columns, start=1):
        width = max(10, min(26, len(str(col)) + 3))
        if "url" in str(col).lower() or "definition" in str(col).lower() or "note" in str(col).lower():
            width = 48
        sh.column_dimensions[get_column_letter(idx)].width = width
    if table_name and len(data) > 0:
        ref = f"A1:{get_column_letter(len(data.columns))}{len(data)+1}"
        tab = Table(displayName=table_name, ref=ref)
        style = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        tab.tableStyleInfo = style
        sh.add_table(tab)
    return sh

cap_sh = add_df_sheet("Depot_Capacities", depot_capacities, "DepotCapacities")
for row in range(2, cap_sh.max_row + 1):
    cap_sh[f"C{row}"].font = Font(color=GREEN_TEXT)
    cap_sh[f"C{row}"].number_format = '#,##0.00'
    cap_sh[f"E{row}"].font = Font(color=GREEN_TEXT, underline="single")
    cap_sh[f"C{row}"].comment = Comment("Official published KPC storage capacity. Source: " + KPC_SOURCE, "OpenAI")

cal_sh = add_df_sheet("Calibration_Anchors", calibration_anchors, "CalibrationAnchors")
for row in range(2, cal_sh.max_row + 1):
    cal_sh[f"B{row}"].number_format = '#,##0.00'
    cal_sh[f"B{row}"].font = Font(color=GREEN_TEXT if cal_sh[f"C{row}"].value == "OFFICIAL_EPRA" else BLUE_TEXT)
    cal_sh[f"D{row}"].font = Font(color=GREEN_TEXT, underline="single")

# Primary data sheets.
ops_sh = add_df_sheet("Daily_Operations", daily_ops, "DailyOperations")
# Format date/numeric columns and conditional highlights without styling every cell heavily.
col_idx = {c: i+1 for i, c in enumerate(daily_ops.columns)}
for row in range(2, ops_sh.max_row + 1):
    ops_sh.cell(row=row, column=col_idx["date"]).number_format = "yyyy-mm-dd"
# Conditional stock status.
status_col = get_column_letter(col_idx["stock_status"])
ops_sh.conditional_formatting.add(f"{status_col}2:{status_col}{ops_sh.max_row}", FormulaRule(formula=[f'${status_col}2="CRITICAL"'], fill=PatternFill("solid", fgColor=ORANGE)))
ops_sh.conditional_formatting.add(f"{status_col}2:{status_col}{ops_sh.max_row}", FormulaRule(formula=[f'${status_col}2="STOCKOUT"'], fill=PatternFill("solid", fgColor=RED)))

batch_sh = add_df_sheet("Pipeline_Batches", auto_batches, "PipelineBatches")
event_sh = add_df_sheet("Equipment_Events", equipment_events, "EquipmentEvents")
roi_data_sh = add_df_sheet("ROI_Comparison_Data", roi_compare, "ROIComparisonData")
cost_sh = add_df_sheet("Cost_Assumptions", cost_parameters, "CostAssumptions")
for row in range(2, cost_sh.max_row + 1):
    cost_sh[f"B{row}"].font = Font(color=BLUE_TEXT)
    cost_sh[f"B{row}"].fill = PatternFill("solid", fgColor="FFF2CC")
    cost_sh[f"B{row}"].number_format = '#,##0.00;[Red](#,##0.00);-'
    cost_sh[f"B{row}"].comment = Comment(str(cost_sh[f"E{row}"].value), "OpenAI")

dict_sh = add_df_sheet("Data_Dictionary", data_dictionary, "DataDictionary")
dict_sh.column_dimensions["A"].width = 34
dict_sh.column_dimensions["B"].width = 100
for row in range(2, dict_sh.max_row + 1):
    dict_sh[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")

# ROI model with formulas.
roi_sh = wb.create_sheet("ROI_Model")
roi_sh.sheet_view.showGridLines = False
roi_sh.merge_cells("A1:K2")
roi_sh["A1"] = "2026 Scenario ROI Model - Baseline vs Autonomous Control Plane"
roi_sh["A1"].font = Font(size=16, bold=True, color=WHITE)
roi_sh["A1"].fill = PatternFill("solid", fgColor=DARK)
roi_sh["A1"].alignment = Alignment(vertical="center")
headers = ["Month", "Baseline Unfulfilled m3", "Autonomous Unfulfilled m3", "Stockout m3 Avoided", "Baseline Emergency m3", "Autonomous Emergency m3", "Emergency m3 Avoided", "Autonomous Batch Actions", "Stockout Value Saved (KES)", "Emergency Logistics Saved (KES)", "Gross Savings (KES)"]
for c, h in enumerate(headers, start=1):
    cell = roi_sh.cell(row=4, column=c, value=h)
    cell.fill = PatternFill("solid", fgColor=HEADER)
    cell.font = Font(bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", wrap_text=True)

# Cost assumptions: rows in Cost_Assumptions are fixed ordering from cost_parameters.
for i, r in roi_compare.iterrows():
    excel_row = 5 + i
    roi_sh.cell(excel_row, 1, r["month_key"])
    roi_sh.cell(excel_row, 2, float(r["baseline_unfulfilled_m3"])).font = Font(color=GREEN_TEXT)
    roi_sh.cell(excel_row, 3, float(r["autonomous_unfulfilled_m3"])).font = Font(color=GREEN_TEXT)
    roi_sh.cell(excel_row, 4, f"=MAX(B{excel_row}-C{excel_row},0)").font = Font(color=BLACK)
    roi_sh.cell(excel_row, 5, float(r["baseline_emergency_m3"])).font = Font(color=GREEN_TEXT)
    roi_sh.cell(excel_row, 6, float(r["autonomous_emergency_m3"])).font = Font(color=GREEN_TEXT)
    roi_sh.cell(excel_row, 7, f"=MAX(E{excel_row}-F{excel_row},0)").font = Font(color=BLACK)
    roi_sh.cell(excel_row, 8, float(r["autonomous_batch_actions"])).font = Font(color=GREEN_TEXT)
    roi_sh.cell(excel_row, 9, f"=D{excel_row}*Cost_Assumptions!$B$2").font = Font(color=BLACK)
    roi_sh.cell(excel_row, 10, f"=G{excel_row}*Cost_Assumptions!$B$3").font = Font(color=BLACK)
    roi_sh.cell(excel_row, 11, f"=I{excel_row}+J{excel_row}").font = Font(color=BLACK)

last_month_row = 4 + len(roi_compare)
total_row = last_month_row + 1
roi_sh.cell(total_row, 1, "2026 TOTAL / SCENARIO")
roi_sh.cell(total_row, 1).font = Font(bold=True)
for c in range(2, 12):
    letter = get_column_letter(c)
    roi_sh.cell(total_row, c, f"=SUM({letter}5:{letter}{last_month_row})")
    roi_sh.cell(total_row, c).font = Font(bold=True, color=BLACK)
    roi_sh.cell(total_row, c).border = Border(top=Side(style="thin", color=BLACK))

sec = total_row + 3
roi_sh[f"A{sec}"] = "Business Case Summary"
roi_sh[f"A{sec}"].font = Font(bold=True, color=WHITE)
roi_sh[f"A{sec}"].fill = PatternFill("solid", fgColor=DARK)
roi_sh.merge_cells(start_row=sec, start_column=1, end_row=sec, end_column=4)
summary_items = [
    ("Gross operational savings (KES)", f"=K{total_row}"),
    ("Implementation cost (KES)", "=Cost_Assumptions!B5"),
    ("Annual support cost (KES)", "=Cost_Assumptions!B6"),
    ("Total year-1 platform cost (KES)", f"=B{sec+2}+B{sec+3}"),
    ("Net year-1 benefit (KES)", f"=B{sec+1}-B{sec+4}"),
    ("Year-1 ROI", f"=IF(B{sec+4}=0,0,B{sec+5}/B{sec+4})"),
    ("Simple payback (months)", f"=IF(B{sec+1}=0,0,B{sec+4}/(B{sec+1}/12))"),
]
for j, (label, formula) in enumerate(summary_items, start=1):
    r = sec + j
    roi_sh[f"A{r}"] = label
    roi_sh[f"B{r}"] = formula
    roi_sh[f"B{r}"].font = Font(color=BLACK)
    if "ROI" in label:
        roi_sh[f"B{r}"].number_format = "0.0%"
    elif "months" in label:
        roi_sh[f"B{r}"].number_format = "0.0x"
    else:
        roi_sh[f"B{r}"].number_format = 'KES #,##0;[Red](KES #,##0);-'
roi_sh.column_dimensions["A"].width = 26
for c in range(2, 12):
    roi_sh.column_dimensions[get_column_letter(c)].width = 22
roi_sh.freeze_panes = "A5"
for row in roi_sh.iter_rows(min_row=5, max_row=total_row, min_col=2, max_col=11):
    for cell in row:
        cell.number_format = '#,##0.00;[Red](#,##0.00);-'

# Monthly gross savings chart.
chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "2026 Monthly Gross Savings (Scenario)"
chart.y_axis.title = "KES"
chart.x_axis.title = "Month"
data_ref = Reference(roi_sh, min_col=11, min_row=4, max_row=last_month_row)
cats_ref = Reference(roi_sh, min_col=1, min_row=5, max_row=last_month_row)
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats_ref)
chart.height = 8
chart.width = 16
roi_sh.add_chart(chart, f"D{sec+1}")

# Validation sheet with formulas.
val_sh = wb.create_sheet("Validation")
val_sh.sheet_view.showGridLines = False
val_sh.merge_cells("A1:D2")
val_sh["A1"] = "Dataset Integrity Checks"
val_sh["A1"].font = Font(size=16, bold=True, color=WHITE)
val_sh["A1"].fill = PatternFill("solid", fgColor=DARK)
checks = [
    ("Capacity breach rows", f'=COUNTIF(Daily_Operations!{get_column_letter(col_idx["capacity_breach_flag"])}:{get_column_letter(col_idx["capacity_breach_flag"])},1)', 0),
    ("Inventory reconciliation errors", f'=COUNTIF(Daily_Operations!{get_column_letter(col_idx["inventory_reconciliation_error"])}:{get_column_letter(col_idx["inventory_reconciliation_error"])},1)', 0),
    ("Negative closing stock rows", f'=COUNTIF(Daily_Operations!{get_column_letter(col_idx["closing_stock_m3"])}:{get_column_letter(col_idx["closing_stock_m3"])},"<0")', 0),
    ("Rows in primary dataset", f'=COUNTA(Daily_Operations!A:A)-1', len(daily_ops)),
    ("Scenario-future rows", f'=COUNTIF(Daily_Operations!{get_column_letter(col_idx["record_status"])}:{get_column_letter(col_idx["record_status"])},"SCENARIO_FUTURE")', int((daily_ops["record_status"]=="SCENARIO_FUTURE").sum())),
]
val_sh["A4"] = "Check"
val_sh["B4"] = "Formula Result"
val_sh["C4"] = "Expected"
val_sh["D4"] = "Status"
for cell in val_sh[4]:
    cell.fill = PatternFill("solid", fgColor=HEADER)
    cell.font = Font(bold=True, color=WHITE)
for i, (label, formula, expected) in enumerate(checks, start=5):
    val_sh[f"A{i}"] = label
    val_sh[f"B{i}"] = formula
    val_sh[f"C{i}"] = expected
    val_sh[f"D{i}"] = f'=IF(B{i}=C{i},"PASS","REVIEW")'
    val_sh[f"B{i}"].font = Font(color=BLACK)
    val_sh[f"C{i}"].font = Font(color=GRAY_TEXT)
    val_sh[f"D{i}"].font = Font(color=PURPLE_TEXT)
val_sh.column_dimensions["A"].width = 36
val_sh.column_dimensions["B"].width = 20
val_sh.column_dimensions["C"].width = 16
val_sh.column_dimensions["D"].width = 14

# Executive summary with formulas referencing Daily_Operations / batches.
exec_sh = wb.create_sheet("Executive_Summary")
exec_sh.sheet_view.showGridLines = False
exec_sh.merge_cells("A1:H2")
exec_sh["A1"] = "Autonomous Supply Chain - Executive Dataset Snapshot"
exec_sh["A1"].font = Font(size=17, bold=True, color=WHITE)
exec_sh["A1"].fill = PatternFill("solid", fgColor=DARK)
exec_sh["A1"].alignment = Alignment(vertical="center")
metrics = [
    ("Rows", f'=COUNTA(Daily_Operations!A:A)-1', "records"),
    ("Depots", 4, "depots"),
    ("Products", 4, "products"),
    ("Autonomous batch triggers", f'=SUM(Daily_Operations!{get_column_letter(col_idx["batch_triggered"])}:{get_column_letter(col_idx["batch_triggered"])})', "actions"),
    ("Emergency replenishments", f'=SUM(Daily_Operations!{get_column_letter(col_idx["emergency_replenishment"])}:{get_column_letter(col_idx["emergency_replenishment"])})', "actions"),
    ("Capacity breaches", f'=SUM(Daily_Operations!{get_column_letter(col_idx["capacity_breach_flag"])}:{get_column_letter(col_idx["capacity_breach_flag"])})', "rows"),
]
for i, (label, value, unit) in enumerate(metrics, start=4):
    exec_sh[f"A{i}"] = label
    exec_sh[f"B{i}"] = value
    exec_sh[f"C{i}"] = unit
    exec_sh[f"A{i}"].font = Font(bold=True, color=GRAY_TEXT)
    exec_sh[f"B{i}"].font = Font(size=13, bold=True, color=TEAL_TEXT if isinstance(value, str) and value.startswith('=') else GRAY_TEXT)
    exec_sh[f"B{i}"].fill = PatternFill("solid", fgColor=TEAL)
exec_sh.column_dimensions["A"].width = 34
exec_sh.column_dimensions["B"].width = 22
exec_sh.column_dimensions["C"].width = 15

# Order sheets.
desired_order = ["README", "Executive_Summary", "Validation", "ROI_Model", "Cost_Assumptions", "Calibration_Anchors", "Depot_Capacities", "Daily_Operations", "Pipeline_Batches", "Equipment_Events", "ROI_Comparison_Data", "Data_Dictionary"]
wb._sheets = [wb[s] for s in desired_order]
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(xlsx_path)

# Basic programmatic QA.
assert daily_ops["capacity_breach_flag"].sum() == 0, "Capacity breach detected"
assert daily_ops["inventory_reconciliation_error"].sum() == 0, "Inventory reconciliation error detected"
assert (daily_ops["closing_stock_m3"] >= -0.01).all(), "Negative stock detected"
assert (daily_ops["closing_stock_m3"] <= daily_ops["tank_capacity_m3"] + 0.01).all(), "Capacity overflow detected"
assert len(daily_ops) == len(base_features), "Unexpected primary row count"
assert set(daily_ops["record_status"].unique()) == {"SIMULATED_HISTORY", "SCENARIO_FUTURE"}

# Summary JSON for quick inspection.
summary = {
    "rows_daily_operations": int(len(daily_ops)),
    "rows_pipeline_batches": int(len(auto_batches)),
    "rows_control_plane_events": int(len(control_plane_events)),
    "capacity_breaches": int(daily_ops["capacity_breach_flag"].sum()),
    "reconciliation_errors": int(daily_ops["inventory_reconciliation_error"].sum()),
    "stockout_rows": int((daily_ops["stock_status"] == "STOCKOUT").sum()),
    "future_scenario_rows": int((daily_ops["record_status"] == "SCENARIO_FUTURE").sum()),
    "seed": SEED,
}
with open(os.path.join(OUT_DIR, "qa_summary.json"), "w", encoding="utf-8") as f:
    json.dump(summary, f, indent=2)

# Zip everything.
zip_path = "/mnt/data/KPC_Synthetic_Control_Plane_Dataset_v1.zip"
with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
    for root, dirs, files in os.walk(OUT_DIR):
        for fname in files:
            fpath = os.path.join(root, fname)
            zf.write(fpath, arcname=os.path.relpath(fpath, OUT_DIR))

print(json.dumps(summary, indent=2))
print(xlsx_path)
print(zip_path)
